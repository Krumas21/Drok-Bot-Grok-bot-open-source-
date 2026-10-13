"""Shell (foreground and background), workspace files, search, and document creation."""

import asyncio
import base64
import csv
import difflib
import io
import json
import shlex
import time
import uuid
from pathlib import Path

from . import Tool, ToolContext, ToolResult, i, params, s

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
TEXT_FORMATS = {"md", "txt", "csv", "json", "yaml", "yml", "py", "js", "ts", "sh"}
PANDOC_FORMATS = {"docx", "pptx", "odt", "epub", "html"}
PDF_STYLE = """<style>
body { font-family: 'DejaVu Sans', 'Liberation Sans', sans-serif; font-size: 11pt; line-height: 1.5; color: #222;
       max-width: 760px; margin: 0 auto; padding: 24px; }
h1, h2, h3 { line-height: 1.2; margin-top: 1.4em; } h1 { font-size: 22pt; } h2 { font-size: 15pt; }
table { border-collapse: collapse; margin: 1em 0; } th, td { border: 1px solid #bbb; padding: 4px 8px; }
th { background: #f2efe9; } code, pre { font-family: 'DejaVu Sans Mono', monospace; font-size: 9pt; }
pre { background: #f5f3ef; padding: 8px; white-space: pre-wrap; } img { max-width: 100%; }
</style>"""


# Every step re-sends earlier tool output, so this is paid again on each later step of the turn. Agents that need all
# of a long output redirect it to a file and read_file it in pages.
def _clip(text: str, head: int = 2500, tail: int = 7500) -> str:
    if len(text) <= head + tail:
        return text
    return (f"{text[:head]}\n[... {len(text) - head - tail} chars cut - to see everything, redirect the output to a "
            f"file and read_file it ...]\n{text[-tail:]}")


def _ws_path(ctx: ToolContext, path: str) -> Path:
    """Map an agent path (/workspace/x or x) to the real workspace dir, refusing escapes."""
    root = ctx.app.settings.workspace
    rel = path.removeprefix("/workspace").lstrip("/") if path.startswith("/workspace") else path.lstrip("/")
    full = (root / rel).resolve()
    if full != root and root not in full.parents:
        raise ValueError("Path is outside /workspace")
    return full


def _container(ctx: ToolContext, real: Path) -> str:
    """The same file as the computer sees it."""
    return "/workspace/" + str(real.relative_to(ctx.app.settings.workspace)) if real != ctx.app.settings.workspace \
        else "/workspace"


async def shell(ctx: ToolContext, command: str, timeout: int = 120, background: bool = False) -> str:
    app = ctx.app
    if background:
        r = await app.computer.bg_start(app.vault.fill(command), screen=ctx.agent["screen"])
        return (f"Started background process #{r['id']} (log {r['log']}). Check it with process(action='status' or "
                f"'log', id={r['id']}); use wake_me to come back to it later.")
    r = await app.computer.exec(app.vault.fill(command), timeout=min(max(timeout, 1), 1800), screen=ctx.agent["screen"])
    return f"exit code: {r['exit_code']}\n{_clip(r['output'])}"


async def process(ctx: ToolContext, action: str = "list", id: str = "", lines: int = 60, timeout: int = 120) -> str:
    c = ctx.app.computer
    if action == "list":
        rows = (await c.bg("list"))["processes"]
        return "\n".join(f"#{p['id']} {'running' if p['running'] else 'exit ' + str(p['exit_code'])} "
                         f"{p['runtime']}s - {p['command']}" for p in rows) or "No background processes."
    if not id:
        return "Pass the process id."
    if action in ("log", "status", "wait"):
        deadline = time.monotonic() + min(max(timeout, 1), 600)
        while True:
            r = await c.bg("log", id=id, lines=min(max(lines, 1), 400))
            if action != "wait" or not r["running"] or time.monotonic() >= deadline:
                break
            await asyncio.sleep(2)
        state = "running" if r["running"] else f"finished (exit {r['exit_code']})"
        return f"#{id} {state}, {r['runtime']}s\n{r['log']}"
    if action == "kill":
        await c.bg("kill", id=id)
        return f"Killed #{id}."
    return "action must be list, status, log, wait or kill."


async def read_file(ctx: ToolContext, path: str, offset: int = 0) -> str | ToolResult:
    p = _ws_path(ctx, path)
    if p.is_dir():
        return "That is a directory; use list_files."
    if p.suffix.lower() in IMAGE_EXT:
        return ToolResult(f"Image {path} ({p.stat().st_size} bytes)", image=base64.b64encode(p.read_bytes()).decode())
    if p.suffix.lower() in (".pdf", ".docx", ".pptx", ".xlsx", ".odt", ".epub"):
        return await _convert_to_text(ctx, p)
    text = p.read_text(errors="replace")
    chunk = text[offset:offset + 30000]
    more = f"\n[... {len(text) - offset - len(chunk)} more chars; read again with offset={offset + len(chunk)}]" \
        if offset + len(chunk) < len(text) else ""
    return chunk + more


async def _convert_to_text(ctx: ToolContext, p: Path) -> str:
    src = shlex.quote(_container(ctx, p))
    if p.suffix.lower() == ".pdf":
        cmd = f"pdftotext -layout {src} - | head -c 60000"
    elif p.suffix.lower() == ".xlsx":
        cmd = (f"python3 -c \"import openpyxl,sys; wb=openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
               f"for ws in wb: print('## '+ws.title); [print(','.join('' if v is None else str(v) for v in r)) "
               f"for r in ws.iter_rows(values_only=True)]\" {src} | head -c 60000")
    else:
        cmd = f"pandoc {src} -t plain | head -c 60000"
    r = await ctx.app.computer.exec(cmd, timeout=120, screen=ctx.agent["screen"])
    return r["output"] if r["exit_code"] == 0 else f"Couldn't extract text (exit {r['exit_code']}): {r['output'][:500]}"


async def write_file(ctx: ToolContext, path: str, content: str, append: bool = False) -> str:
    p = _ws_path(ctx, path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a" if append else "w") as f:
        f.write(content)
    return f"{'Appended' if append else 'Wrote'} {len(content)} chars to {_container(ctx, p)}"


async def edit_file(ctx: ToolContext, path: str, old_string: str, new_string: str, replace_all: bool = False) -> str:
    p = _ws_path(ctx, path)
    if not p.is_file():
        return f"Not found: {path}"
    text = p.read_text()
    count = text.count(old_string) if old_string else 0
    if count == 0:
        near = difflib.get_close_matches(old_string.strip().splitlines()[0] if old_string.strip() else "",
                                         text.splitlines(), n=3, cutoff=0.6)
        return "Error: old_string not found." + (" Closest lines:\n" + "\n".join(near) if near else "")
    if count > 1 and not replace_all:
        return f"Error: old_string appears {count} times. Add surrounding lines to make it unique, or replace_all=true."
    updated = text.replace(old_string, new_string) if replace_all else text.replace(old_string, new_string, 1)
    p.write_text(updated)
    diff = "".join(difflib.unified_diff(text.splitlines(True), updated.splitlines(True), "before", "after", n=2))
    return f"Edited {_container(ctx, p)} ({count if replace_all else 1} replacement).\n{diff[:4000]}"


async def list_files(ctx: ToolContext, path: str = "/workspace") -> str:
    p = _ws_path(ctx, path)
    if not p.exists():
        return "Not found."
    rows = []
    for child in sorted(p.iterdir())[:500]:
        if child.name.startswith("."):
            continue
        rows.append(f"{child.name}/" if child.is_dir() else f"{child.name}  ({child.stat().st_size} B)")
    return "\n".join(rows) or "(empty)"


async def search_files(ctx: ToolContext, pattern: str, path: str = "/workspace", glob: str = "",
                       names_only: bool = False) -> str:
    where = shlex.quote(_container(ctx, _ws_path(ctx, path)))
    if names_only:
        cmd = f"rg --files {where} 2>/dev/null | rg -i -- {shlex.quote(pattern)} | head -200"
    else:
        g = f"-g {shlex.quote(glob)} " if glob else ""
        cmd = (f"rg -n -S --max-columns 300 --max-count 20 {g}-- {shlex.quote(pattern)} {where} 2>/dev/null "
               f"| head -300")
    r = await ctx.app.computer.exec(cmd, timeout=60, screen=ctx.agent["screen"])
    return r["output"].strip() or "No matches."


async def make_file(ctx: ToolContext, path: str, content: str, title: str = "", format: str = "") -> str:
    """Write a deliverable. Markdown in -> docx/pptx/odt/epub/html/pdf out; CSV or JSON rows in -> xlsx."""
    app = ctx.app
    out = _ws_path(ctx, path)
    fmt = (format or out.suffix.lstrip(".")).lower()
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt in TEXT_FORMATS or (fmt == "html" and content.lstrip().lower().startswith(("<!doctype", "<html"))):
        out.write_text(content)
        return f"Created {_container(ctx, out)} ({out.stat().st_size} bytes)."
    tmp = app.settings.workspace / ".drok" / "tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    job = uuid.uuid4().hex[:10]
    q_out, q_title = shlex.quote(_container(ctx, out)), shlex.quote(title or out.stem)
    if fmt in PANDOC_FORMATS | {"pdf"}:
        src = tmp / f"{job}.md"
        src.write_text(content)
        q_src = shlex.quote(_container(ctx, src))
        if fmt == "pdf":
            style = tmp / "pdf-style.html"
            style.write_text(PDF_STYLE)
            html = f"/tmp/{job}.html"
            cmd = (f"pandoc {q_src} -s --metadata pagetitle={q_title} -H {shlex.quote(_container(ctx, style))} "
                   f"-o {html} && chromium --headless=new --no-sandbox --disable-gpu --no-pdf-header-footer "
                   f"--user-data-dir=/tmp/drok-pdf-profile --print-to-pdf={q_out} file://{html} 2>/dev/null")
        else:
            cmd = f"pandoc {q_src} -s --metadata title={q_title} -o {q_out}"
    elif fmt == "xlsx":
        src = tmp / f"{job}.json"
        try:
            data = json.loads(content)
        except ValueError:  # treat as CSV
            data = {"Sheet1": list(csv.reader(io.StringIO(content)))}
        if isinstance(data, list):
            data = {"Sheet1": data}
        src.write_text(json.dumps(data))
        cmd = f"python3 - {shlex.quote(_container(ctx, src))} {q_out} <<'PY'\n{XLSX_SCRIPT}\nPY"
    else:
        return (f"Unsupported format '{fmt}'. Use md, txt, csv, json, html, docx, pptx, odt, epub, pdf or xlsx - or "
                "write a script with shell for anything else.")
    r = await app.computer.exec(cmd, timeout=180, screen=ctx.agent["screen"])
    for f in tmp.glob(f"{job}.*"):
        f.unlink()
    if r["exit_code"] != 0 or not out.exists():
        return f"Failed to create {fmt} (exit {r['exit_code']}): {r['output'][-1500:]}"
    return f"Created {_container(ctx, out)} ({out.stat().st_size} bytes)."


XLSX_SCRIPT = """import json, sys
import openpyxl
from openpyxl.styles import Font
data = json.load(open(sys.argv[1]))
wb = openpyxl.Workbook()
wb.remove(wb.active)
for name, rows in data.items():
    ws = wb.create_sheet(str(name)[:31])
    if rows and isinstance(rows[0], dict):
        keys = list(rows[0].keys())
        rows = [keys] + [[r.get(k) for k in keys] for r in rows]
    for row in rows:
        ws.append(list(row))
    for cell in ws[1] if rows else []:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
wb.save(sys.argv[2])"""


def tools(app) -> list[Tool]:
    return [
        Tool("shell", "Run a bash command on the shared Linux computer (cwd /workspace). Use for git, builds, "
             "scripts, package installs (sudo works), and bulk work via scripts. background=true for long-running "
             "jobs (servers, long scripts) - then use process. Reference stored secrets as {{secret:NAME}}.",
             params(["command"], command=s("Bash command"), timeout=i("Seconds before kill, default 120, max 1800"),
                    background={"type": "boolean", "description": "Run detached and return immediately"}),
             shell, risk="risky", subject=lambda a: a.get("command", ""), timeout=1830),
        Tool("process", "Manage background processes started with shell(background=true): list, status/log "
             "(tail of output), wait (block until it exits or `timeout` seconds pass, max 600), kill. For waits "
             "longer than 10 minutes use wake_me instead.",
             params([], action=s("list|status|log|wait|kill"), id=s("Process id"), lines=i("Log lines, default 60"),
                    timeout=i("Seconds to wait (action=wait), default 120")),
             process, risk="safe", timeout=630, subject=lambda a: f"{a.get('action', 'list')} {a.get('id', '')}"),
        Tool("read_file", "Read a file from /workspace: text (with offset paging), images (viewed), or PDF/DOCX/PPTX/"
             "XLSX (converted to text).",
             params(["path"], path=s("Path, e.g. /workspace/notes.md"), offset=i("Char offset for long files")),
             read_file, risk="safe"),
        Tool("write_file", "Create or overwrite a text file in /workspace (append=true to append).",
             params(["path", "content"], path=s("Path inside /workspace"), content=s("Full file content"),
                    append={"type": "boolean"}),
             write_file, risk="safe", subject=lambda a: a.get("path", "")),
        Tool("edit_file", "Replace an exact snippet in a file (must be unique unless replace_all). Returns a diff. "
             "Prefer this over rewriting whole files.",
             params(["path", "old_string", "new_string"], path=s("File path"), old_string=s("Exact current text"),
                    new_string=s("Replacement"), replace_all={"type": "boolean"}),
             edit_file, risk="safe", subject=lambda a: a.get("path", "")),
        Tool("list_files", "List a directory in /workspace.", params([], path=s("Directory, default /workspace")),
             list_files, risk="safe"),
        Tool("search_files", "Search file contents with ripgrep (regex), or find files by name (names_only=true).",
             params(["pattern"], pattern=s("Regex or name fragment"), path=s("Where, default /workspace"),
                    glob=s("Filter, e.g. *.py"), names_only={"type": "boolean"}),
             search_files, risk="safe", subject=lambda a: a.get("pattern", "")),
        Tool("make_file", "Create a polished deliverable in /workspace. Markdown content -> docx, pdf, pptx (each # "
             "or ## heading starts a slide), odt, epub or html; CSV text or JSON rows ({sheet: rows} for several "
             "sheets) -> xlsx; or plain md/txt/csv/json. Format comes from the extension unless given.",
             params(["path", "content"], path=s("Output path, e.g. /workspace/reports/q3.pdf"),
                    content=s("Markdown (documents) or CSV/JSON rows (spreadsheets)"), title=s("Document title"),
                    format=s("Override: docx|pdf|pptx|xlsx|html|odt|epub|md|txt|csv|json")),
             make_file, risk="safe", timeout=200, subject=lambda a: a.get("path", "")),
    ]
